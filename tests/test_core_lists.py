import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import docx  # noqa: E402

from core.document_model import Document, Paragraph, Run  # noqa: E402
from core.docx_io import load_docx, save_docx  # noqa: E402


def _p(text, list_type=None, level=0):
    return Paragraph(runs=[Run(text)], list_type=list_type, list_level=level)


def test_list_labels_numbering_and_restart():
    doc = Document(paragraphs=[
        _p("a", "number"), _p("b", "number"), _p("b1", "number", 1), _p("b2", "number", 1),
        _p("c", "number"), _p("c1", "number", 1), _p("bullet", "bullet"), _p("deep", "bullet", 2),
        _p("pause"), _p("neu", "number"),
    ])
    assert doc.list_labels() == ["1.", "2.", "a.", "b.", "3.", "a.", "•", "▪", None, "1."]


def test_roman_and_letters():
    doc = Document(paragraphs=[_p("x", "number", 2) for _ in range(4)])
    assert doc.list_labels() == ["i.", "ii.", "iii.", "iv."]
    doc = Document(paragraphs=[_p("x", "number", 1) for _ in range(27)])
    assert doc.list_labels()[-2:] == ["z.", "aa."]


def test_lists_roundtrip_docx(tmp_path):
    doc = Document(paragraphs=[
        _p("Einkauf"), _p("Milch", "bullet"), _p("Bio", "bullet", 1),
        _p("Schritt", "number"), _p("Unterschritt", "number", 2),
    ])
    path = str(tmp_path / "listen.docx")
    save_docx(doc, path)
    raw = docx.Document(path)
    assert [p.style.name for p in raw.paragraphs] == [
        "Normal", "List Bullet", "List Bullet 2", "List Number", "List Number 3"]
    loaded = load_docx(path)
    assert [(p.list_type, p.list_level, p.style) for p in loaded.paragraphs] == [
        (None, 0, "Normal"), ("bullet", 0, "Normal"), ("bullet", 1, "Normal"),
        ("number", 0, "Normal"), ("number", 2, "Normal")]


def test_direct_numbering_is_detected(tmp_path):
    import copy
    raw = docx.Document()
    bullet_numpr = raw.styles["List Bullet"].element.pPr.numPr
    para = raw.add_paragraph("direkt", style="List Paragraph")
    para._p.get_or_add_pPr().append(copy.deepcopy(bullet_numpr))
    raw.add_paragraph("normal")
    path = str(tmp_path / "direkt.docx")
    raw.save(path)
    loaded = load_docx(path)
    assert loaded.paragraphs[0].list_type == "bullet"
    assert loaded.paragraphs[1].list_type is None
