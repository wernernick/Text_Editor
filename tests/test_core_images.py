import os
import struct
import sys
import zlib

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import docx  # noqa: E402
import pytest  # noqa: E402

from core.document_model import Document, Image, Paragraph, Run  # noqa: E402
from core.docx_io import load_docx, save_docx  # noqa: E402


def _png(width, height):
    """Minimales PNG (Graustufen) ohne Zusatzbibliotheken."""
    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))
    raw = b"".join(b"\x00" + b"\x80" * width for _ in range(height))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 0, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def test_image_from_file_scales_down(tmp_path):
    path = tmp_path / "logo.png"
    path.write_bytes(_png(200, 100))
    img = Image.from_file(str(path), max_width_pt=100)
    assert img.filename == "logo.png"
    assert img.native_size_pt() == (200, 100)
    assert img.display_size_pt() == (100, 50)
    assert Image.from_file(str(path)).width_pt is None


def test_unknown_image_format():
    with pytest.raises(ValueError):
        Image(data=b"kein bild", filename="x.txt").native_size_pt()


def test_image_roundtrip(tmp_path):
    doc = Document.from_text("Text davor")
    doc.insert_image(1, Image(data=_png(40, 20), width_pt=80))
    doc.paragraphs.append(Paragraph(runs=[Run("danach")]))
    path = str(tmp_path / "bild.docx")
    save_docx(doc, path)
    assert len(docx.Document(path).inline_shapes) == 1

    loaded = load_docx(path)
    assert [p.image is not None for p in loaded.paragraphs] == [False, True, False]
    img = loaded.paragraphs[1].image
    assert img.data == _png(40, 20)
    assert img.display_size_pt() == pytest.approx((80, 40))
    assert loaded.paragraphs[1].alignment == "center"
    assert loaded.list_labels() == [None, None, None]


def test_text_and_image_in_one_paragraph_are_split(tmp_path):
    import io
    raw = docx.Document()
    para = raw.add_paragraph("Bildunterschrift ")
    para.add_run().add_picture(io.BytesIO(_png(10, 10)))
    path = str(tmp_path / "gemischt.docx")
    raw.save(path)
    loaded = load_docx(path)
    assert loaded.paragraphs[0].text == "Bildunterschrift "
    assert loaded.paragraphs[1].image is not None


def test_unsupported_image_format_does_not_break_saving(tmp_path):
    emf = Image(data=b"\x01\x00\x00\x00" + b"\x00" * 60, filename="image1.emf", width_pt=100, height_pt=50)
    assert not emf.is_supported
    assert emf.display_size_pt() == (100, 50)
    doc = Document.from_text("Vorher")
    doc.insert_image(1, emf)
    doc.paragraphs.append(Paragraph(runs=[Run("Nachher")]))
    path = str(tmp_path / "emf.docx")
    save_docx(doc, path)
    loaded = load_docx(path)
    assert [p.text for p in loaded.paragraphs] == ["Vorher", "[Bild nicht unterstützt: image1.emf]", "Nachher"]


def test_loaded_emf_from_word_saves(tmp_path):
    """Word-Datei mit EMF-Bild: Laden und erneutes Speichern darf nicht abbrechen."""
    import io
    raw = docx.Document()
    raw.add_paragraph().add_run().add_picture(io.BytesIO(_png(10, 10)))
    path = str(tmp_path / "quelle.docx")
    raw.save(path)
    doc = load_docx(path)
    doc.paragraphs[0].image.data = b"\x01\x00\x00\x00" + b"\x00" * 60  # wie ein EMF-Blob
    doc.paragraphs[0].image.filename = "image1.emf"
    save_docx(doc, str(tmp_path / "ziel.docx"))
    assert "nicht unterstützt" in load_docx(str(tmp_path / "ziel.docx")).plain_text()
