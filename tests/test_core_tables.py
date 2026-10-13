import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import docx  # noqa: E402
import pytest  # noqa: E402

from core.document_model import Document, Paragraph, Run, Table  # noqa: E402
from core.docx_io import load_docx, save_docx  # noqa: E402


def test_table_editing():
    table = Table.create(2, 3)
    assert (table.row_count, table.col_count) == (2, 3)
    table.insert_row(1)
    table.insert_column(0)
    assert (table.row_count, table.col_count) == (3, 4)
    table.delete_row(0)
    table.delete_column(3)
    assert (table.row_count, table.col_count) == (2, 3)
    single = Table.create(1, 1)
    single.delete_row(0)
    single.delete_column(0)
    assert (single.row_count, single.col_count) == (1, 1)
    with pytest.raises(ValueError):
        Table.create(0, 2)


def test_document_insert_table_and_plain_text():
    doc = Document.from_text("Vorher\nNachher")
    table = doc.insert_table(1, 2, 2)
    table.rows[0][0].paragraphs[0].runs = [Run("A")]
    table.rows[1][1].paragraphs[0].runs = [Run("D")]
    assert doc.plain_text() == "Vorher\nA\t\n\tD\nNachher"
    assert doc.list_labels() == [None, None, None]


def test_table_roundtrip_keeps_order_and_formatting(tmp_path):
    table = Table.from_texts([["Name", "Wert"], ["Pi", "3,14"]], header_row=True)
    table.rows[1][0].paragraphs[0].runs[0].bold = True
    table.rows[1][1] = table.rows[1][1].from_text("Zeile 1\nZeile 2")
    doc = Document(paragraphs=[
        Paragraph(runs=[Run("Davor")]),
        Paragraph(table=table),
        Paragraph(runs=[Run("Danach")], alignment="center"),
    ])
    path = str(tmp_path / "tabelle.docx")
    save_docx(doc, path)
    assert len(docx.Document(path).tables) == 1

    loaded = load_docx(path)
    assert [p.table is not None for p in loaded.paragraphs] == [False, True, False]
    out = loaded.paragraphs[1].table
    assert out.texts() == [["Name", "Wert"], ["Pi", "Zeile 1\nZeile 2"]]
    assert out.header_row
    assert out.rows[1][0].paragraphs[0].runs[0].bold
    assert loaded.paragraphs[2].alignment == "center"
