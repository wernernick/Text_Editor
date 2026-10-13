"""Lesen und Schreiben von .docx-Dateien über python-docx."""
from __future__ import annotations

import os

import docx
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from core.document_model import LIST_TYPES, MAX_LIST_LEVEL, Document, Paragraph, Run

_ALIGN_TO_DOCX = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
_DOCX_TO_ALIGN = {value: key for key, value in _ALIGN_TO_DOCX.items()}
_LIST_STYLE_BASE = {"bullet": "List Bullet", "number": "List Number"}


def _list_style_name(list_type: str, level: int) -> str:
    base = _LIST_STYLE_BASE[list_type]
    return base if level <= 0 else f"{base} {level + 1}"


def _parse_list_style(style: str) -> tuple[str, int] | None:
    """"List Bullet 2" -> ("bullet", 1); andere Formatvorlagen -> None."""
    for list_type, base in _LIST_STYLE_BASE.items():
        if style == base:
            return list_type, 0
        if style.startswith(base + " ") and style[len(base) + 1:].isdigit():
            return list_type, min(MAX_LIST_LEVEL, int(style[len(base) + 1:]) - 1)
    return None


def _numbering_list(paragraph) -> tuple[str, int] | None:
    """Erkennt direkte Nummerierung (w:numPr) von Listen, die nicht über Formatvorlagen laufen."""
    p_pr = paragraph._p.pPr
    num_pr = p_pr.numPr if p_pr is not None else None
    if num_pr is None or num_pr.numId is None or num_pr.numId.val == 0:
        return None
    level = num_pr.ilvl.val if num_pr.ilvl is not None else 0
    list_type = "number"
    try:
        numbering = paragraph.part.numbering_part.element
        num = numbering.num_having_numId(num_pr.numId.val)
        abstract_id = num.abstractNumId.val
        for abstract in numbering.findall(qn("w:abstractNum")):
            if abstract.get(qn("w:abstractNumId")) != str(abstract_id):
                continue
            for lvl in abstract.findall(qn("w:lvl")):
                if lvl.get(qn("w:ilvl")) == str(level):
                    fmt = lvl.find(qn("w:numFmt"))
                    if fmt is not None and fmt.get(qn("w:val")) == "bullet":
                        list_type = "bullet"
    except (AttributeError, KeyError, NotImplementedError):
        pass
    return list_type, min(MAX_LIST_LEVEL, level)


def _style_name(paragraph) -> str:
    try:
        return paragraph.style.name if paragraph.style is not None else "Normal"
    except (KeyError, ValueError):
        return "Normal"


def _read_color(run) -> str | None:
    try:
        rgb = run.font.color.rgb if run.font.color is not None and run.font.color.type is not None else None
    except (AttributeError, ValueError):
        return None
    return f"#{rgb}" if rgb is not None else None


def load_docx(path: str) -> Document:
    """Lädt eine .docx-Datei in das Dokumentmodell."""
    source = docx.Document(path)
    paragraphs: list[Paragraph] = []
    for para in source.paragraphs:
        runs = [
            Run(
                text=run.text,
                bold=bool(run.bold),
                italic=bool(run.italic),
                underline=bool(run.underline),
                font_family=run.font.name,
                font_size=run.font.size.pt if run.font.size is not None else None,
                color=_read_color(run),
            )
            for run in para.runs
        ]
        style = _style_name(para)
        list_info = _parse_list_style(style) or _numbering_list(para)
        if list_info and style.startswith(("List Bullet", "List Number")):
            style = "Normal"
        model_para = Paragraph(
            runs=runs,
            alignment=_DOCX_TO_ALIGN.get(para.alignment, "left"),
            style=style,
            list_type=list_info[0] if list_info else None,
            list_level=list_info[1] if list_info else 0,
        )
        model_para.merge_runs()
        paragraphs.append(model_para)
    if not paragraphs:
        paragraphs.append(Paragraph(runs=[]))
    return Document(paragraphs=paragraphs, path=os.path.abspath(path))


def save_docx(doc: Document, path: str) -> None:
    """Speichert das Dokumentmodell als .docx-Datei."""
    target = docx.Document()
    style_names = {style.name for style in target.styles}
    for para in doc.paragraphs:
        out = target.add_paragraph()
        style = para.style
        if para.list_type in LIST_TYPES:
            style = _list_style_name(para.list_type, max(0, min(MAX_LIST_LEVEL, para.list_level)))
        if style in style_names:
            out.style = target.styles[style]
        out.alignment = _ALIGN_TO_DOCX.get(para.alignment, WD_ALIGN_PARAGRAPH.LEFT)
        for run in para.runs:
            out_run = out.add_run(run.text)
            out_run.bold = run.bold
            out_run.italic = run.italic
            out_run.underline = run.underline
            if run.font_family:
                out_run.font.name = run.font_family
            if run.font_size:
                out_run.font.size = Pt(run.font_size)
            if run.color:
                out_run.font.color.rgb = RGBColor.from_string(run.color.lstrip("#").upper())
    folder = os.path.dirname(os.path.abspath(path))
    os.makedirs(folder, exist_ok=True)
    target.save(path)
    doc.path = os.path.abspath(path)
