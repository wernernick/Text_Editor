"""Datenmodell eines Dokuments: Dokument -> Absätze -> Runs (gleich formatierter Text)."""
from __future__ import annotations

import os
from dataclasses import dataclass, field
from typing import Optional

ALIGNMENTS = ("left", "center", "right", "justify")
LIST_TYPES = ("bullet", "number")
MAX_LIST_LEVEL = 2  # Ebenen 0..2, wie die Word-Standardvorlagen "List Bullet", "List Bullet 2", "List Bullet 3"
BULLET_SYMBOLS = ("•", "◦", "▪")


@dataclass
class Run:
    """Ein Textstück mit einheitlicher Formatierung."""

    text: str
    bold: bool = False
    italic: bool = False
    underline: bool = False
    font_family: Optional[str] = None
    font_size: Optional[float] = None  # in Punkt
    color: Optional[str] = None  # "#RRGGBB"

    def same_format(self, other: "Run") -> bool:
        return (
            self.bold == other.bold
            and self.italic == other.italic
            and self.underline == other.underline
            and self.font_family == other.font_family
            and self.font_size == other.font_size
            and self.color == other.color
        )


@dataclass
class Paragraph:
    """Ein Absatz aus mehreren Runs."""

    runs: list[Run] = field(default_factory=list)
    alignment: str = "left"  # left | center | right | justify
    style: str = "Normal"  # z.B. "Heading 1"
    list_type: Optional[str] = None  # None | "bullet" | "number"
    list_level: int = 0  # Einrückungsebene der Liste, 0..MAX_LIST_LEVEL
    table: Optional["Table"] = None  # gesetzt = dieser Block ist eine Tabelle, runs werden ignoriert
    image: Optional["Image"] = None  # gesetzt = dieser Block ist ein Bild, runs werden ignoriert

    @property
    def text(self) -> str:
        return "".join(run.text for run in self.runs)

    def merge_runs(self) -> None:
        """Fasst benachbarte Runs mit gleicher Formatierung zusammen und entfernt leere."""
        merged: list[Run] = []
        for run in self.runs:
            if not run.text:
                continue
            if merged and merged[-1].same_format(run):
                merged[-1].text += run.text
            else:
                merged.append(Run(**vars(run)))
        self.runs = merged


@dataclass
class Image:
    """Ein Bild als eigener Block (Bytes im Dokument, damit nichts von externen Dateien abhängt)."""

    data: bytes
    filename: str = "bild.png"
    width_pt: Optional[float] = None  # None = Originalgröße
    height_pt: Optional[float] = None  # None = proportional zur Breite

    @classmethod
    def from_file(cls, path: str, max_width_pt: Optional[float] = None) -> "Image":
        """Lädt ein Bild (PNG, JPEG, GIF, BMP, TIFF); zu breite Bilder werden proportional verkleinert."""
        with open(path, "rb") as handle:
            data = handle.read()
        image = cls(data=data, filename=os.path.basename(path))
        width, height = image.native_size_pt()
        if max_width_pt and width > max_width_pt:
            image.width_pt = max_width_pt
            image.height_pt = height * max_width_pt / width
        return image

    def native_size_pt(self) -> tuple[float, float]:
        """Originalgröße in Punkt (aus Pixeln und DPI der Datei). Wirft ValueError bei unbekanntem Format."""
        from docx.image.image import Image as DocxImage  # python-docx kennt die Bildformate bereits
        from docx.image.exceptions import UnrecognizedImageError

        try:
            info = DocxImage.from_blob(self.data)
        except UnrecognizedImageError as exc:
            raise ValueError(f"Unbekanntes Bildformat: {self.filename}") from exc
        return info.px_width * 72.0 / (info.horz_dpi or 72), info.px_height * 72.0 / (info.vert_dpi or 72)

    def display_size_pt(self) -> tuple[float, float]:
        """Größe, in der das Bild angezeigt und gespeichert wird."""
        native_w, native_h = self.native_size_pt()
        width = self.width_pt or native_w
        if self.height_pt:
            return width, self.height_pt
        return width, native_h * width / native_w if native_w else native_h


@dataclass
class TableCell:
    """Eine Tabellenzelle; enthält wie in Word eigene Absätze."""

    paragraphs: list[Paragraph] = field(default_factory=lambda: [Paragraph()])

    @property
    def text(self) -> str:
        return "\n".join(p.text for p in self.paragraphs)

    @classmethod
    def from_text(cls, text: str) -> "TableCell":
        return cls(paragraphs=[Paragraph(runs=[Run(line)] if line else []) for line in text.split("\n")])


@dataclass
class Table:
    """Einfache Tabelle: rechteckiges Raster aus Zellen (keine verbundenen Zellen)."""

    rows: list[list[TableCell]] = field(default_factory=list)
    header_row: bool = False  # erste Zeile als Kopfzeile hervorheben

    @classmethod
    def create(cls, row_count: int, col_count: int) -> "Table":
        if row_count < 1 or col_count < 1:
            raise ValueError("Tabelle braucht mindestens 1 Zeile und 1 Spalte")
        return cls(rows=[[TableCell() for _ in range(col_count)] for _ in range(row_count)])

    @classmethod
    def from_texts(cls, rows: list[list[str]], header_row: bool = False) -> "Table":
        width = max((len(r) for r in rows), default=0)
        table = cls.create(max(1, len(rows)), max(1, width))
        for r, row in enumerate(rows):
            for c, text in enumerate(row):
                table.rows[r][c] = TableCell.from_text(text)
        table.header_row = header_row
        return table

    @property
    def row_count(self) -> int:
        return len(self.rows)

    @property
    def col_count(self) -> int:
        return len(self.rows[0]) if self.rows else 0

    def cell(self, row: int, col: int) -> TableCell:
        return self.rows[row][col]

    def texts(self) -> list[list[str]]:
        return [[cell.text for cell in row] for row in self.rows]

    def insert_row(self, index: int | None = None) -> None:
        index = self.row_count if index is None else index
        self.rows.insert(index, [TableCell() for _ in range(max(1, self.col_count))])

    def insert_column(self, index: int | None = None) -> None:
        index = self.col_count if index is None else index
        for row in self.rows:
            row.insert(index, TableCell())

    def delete_row(self, index: int) -> None:
        if self.row_count > 1:
            del self.rows[index]

    def delete_column(self, index: int) -> None:
        if self.col_count > 1:
            for row in self.rows:
                del row[index]


@dataclass
class Document:
    """Ein komplettes Dokument."""

    paragraphs: list[Paragraph] = field(default_factory=list)
    path: Optional[str] = None

    @classmethod
    def new(cls) -> "Document":
        return cls(paragraphs=[Paragraph(runs=[])])

    @classmethod
    def from_text(cls, text: str) -> "Document":
        lines = text.split("\n")
        return cls(paragraphs=[Paragraph(runs=[Run(line)] if line else []) for line in lines])

    def plain_text(self) -> str:
        """Gesamter Text; Tabellenzellen durch Tab, Tabellenzeilen durch Zeilenumbruch getrennt."""
        lines = []
        for p in self.paragraphs:
            if p.table is not None:
                lines.extend("\t".join(c.text.replace("\n", " ") for c in row) for row in p.table.rows)
            else:
                lines.append(p.text)
        return "\n".join(lines)

    def insert_table(self, index: int, row_count: int, col_count: int) -> Table:
        """Fügt an Position index einen Tabellen-Block ein und gibt die Tabelle zurück."""
        table = Table.create(row_count, col_count)
        self.paragraphs.insert(index, Paragraph(table=table))
        return table

    def insert_image(self, index: int, image: Image, alignment: str = "center") -> Image:
        """Fügt an Position index einen Bild-Block ein."""
        self.paragraphs.insert(index, Paragraph(image=image, alignment=alignment))
        return image

    def list_labels(self) -> list[Optional[str]]:
        """Aufzählungszeichen bzw. Nummer je Absatz ("•", "1.", "a." ...), None ohne Liste.

        Nummerierung zählt pro Ebene fortlaufend, bis ein Absatz ohne Liste sie unterbricht;
        eine tiefere Ebene beginnt nach jedem übergeordneten Eintrag neu (wie in Word).
        """
        labels: list[Optional[str]] = []
        counters = [0] * (MAX_LIST_LEVEL + 1)
        for para in self.paragraphs:
            if para.list_type not in LIST_TYPES or para.table is not None or para.image is not None:
                counters = [0] * (MAX_LIST_LEVEL + 1)
                labels.append(None)
                continue
            level = max(0, min(MAX_LIST_LEVEL, para.list_level))
            for deeper in range(level + 1, MAX_LIST_LEVEL + 1):
                counters[deeper] = 0
            if para.list_type == "bullet":
                labels.append(BULLET_SYMBOLS[level])
                continue
            counters[level] += 1
            labels.append(_number_label(counters[level], level))
        return labels

    def word_count(self) -> int:
        return len(self.plain_text().split())

    def char_count(self) -> int:
        return sum(len(p.text) for p in self.paragraphs)


def _number_label(n: int, level: int) -> str:
    """1. / a. / i. je nach Ebene."""
    if level == 1:
        letters = ""
        while n > 0:
            n, rest = divmod(n - 1, 26)
            letters = chr(ord("a") + rest) + letters
        return f"{letters}."
    if level == 2:
        return f"{_roman(n)}."
    return f"{n}."


def _roman(n: int) -> str:
    values = ((1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"), (90, "xc"),
              (50, "l"), (40, "xl"), (10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i"))
    out = ""
    for value, symbol in values:
        while n >= value:
            out += symbol
            n -= value
    return out
